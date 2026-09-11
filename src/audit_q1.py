from pathlib import Path
import csv, json, math
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

def read_csv(name):
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def audit_schedule(rows, purchase_key):
    f=lambda key: np.array([float(r[key]) for r in rows])
    price=f("电价(元/kWh)"); load=f("负载电量(kWh)"); pv=f("光伏电量(kWh)")
    buy=f(purchase_key); ch=f("充电量(kWh)"); dis=f("放电量(kWh)"); spill=f("弃光量(kWh)")
    e0=f("期初储电量(kWh)"); e1=f("期末储电量(kWh)")
    return {
        "rows":len(rows), "first_label":rows[0]["时间段"], "last_label":rows[-1]["时间段"],
        "cost":float(price@buy), "purchase":float(buy.sum()), "charge":float(ch.sum()), "discharge":float(dis.sum()),
        "max_balance_error":float(np.abs(buy+pv+dis-load-ch-spill).max()),
        "max_state_error":float(np.abs(e1-e0-.9*ch+dis/.9).max()),
        "max_continuity_error":float(np.abs(e1[:-1]-e0[1:]).max()),
        "energy_min":float(min(e0.min(),e1.min())), "energy_max":float(max(e0.max(),e1.max())),
        "initial":float(e0[0]), "final":float(e1[-1]),
        "max_charge":float(ch.max()), "max_discharge":float(dis.max()),
        "simultaneous":int(((ch>1e-7)&(dis>1e-7)).sum()), "spill_exceeds_pv":int((spill>pv+1e-7).sum()),
    }

milp_rows=read_csv("outputs/q1_milp_schedule.csv")
dp_rows=read_csv("outputs/q1_dp/dp_schedule_finest_grid.csv")
milp_summary=json.loads((ROOT/"outputs/q1_milp_summary.json").read_text(encoding="utf-8"))
milp=audit_schedule(milp_rows,"计划购电量(kWh)")
dp=audit_schedule(dp_rows,"购电量(kWh)")
raw=list(openpyxl.load_workbook(ROOT/'materials/附件1.xlsx',read_only=True,data_only=True).active.values)[1:]
ordered=raw[-1:]+raw[:-1]
for schedule in (milp_rows,dp_rows):
    for r,source in zip(schedule,ordered):
        assert float(r['电价(元/kWh)']) == source[1]
        assert abs(float(r['负载电量(kWh)'])-source[2]/6)<1e-10
        assert abs(float(r['光伏电量(kWh)'])-source[3]/6)<1e-10
for result in (milp,dp):
    assert result['energy_min']>=1200-1e-6 and result['energy_max']<=10800+1e-6
    assert abs(result['initial']-6000)<1e-6 and abs(result['final']-6000)<1e-6
    assert result['max_charge']<=5000/6+1e-6 and result['max_discharge']<=5000/6+1e-6
    assert result['max_continuity_error']<1e-6

wb=openpyxl.load_workbook(ROOT/"materials/result1.xlsx",read_only=True,data_only=True)
s=wb["计划购电量"]; b=wb["充放电量"]
book_buy=np.array([float(s.cell(i,2).value) for i in range(2,146)])
book_labels=[str(s.cell(i,1).value) for i in range(2,146)]
def template_clock(total_minutes):
    day = total_minutes // 1440
    minute = total_minutes % 1440
    label = f"{minute // 60}:{minute % 60:02d}"
    return label + ("+1" if day else "")

expected_labels=[f"{template_clock(v)}-{template_clock(v+10)}" for v in range(10,1450,10)]
blocks=np.array([[float(b.cell(i,j).value) for j in (2,3)] for i in range(2,8)])
expected_blocks=np.array([[sum(float(milp_rows[k][key]) for k in range(q*24,(q+1)*24)) for key in ("充电量(kWh)","放电量(kWh)")] for q in range(6)])

conv=read_csv("outputs/q1_dp/dp_convergence.csv")
steps=np.array([float(r["state_step_kwh"]) for r in conv])
costs=np.array([float(r["total_cost_yuan"]) for r in conv])

report={"milp":milp,"dp_5kwh":dp,"cross_model":{"dp_cost_gap":dp["cost"]-milp["cost"],"relative_gap":dp["cost"]/milp["cost"]-1},
"result1":{"purchase_max_abs_diff":float(np.abs(book_buy-np.array([float(r['计划购电量(kWh)']) for r in milp_rows[1:]+milp_rows[:1]])).max()),"block_max_abs_diff":float(np.abs(blocks-expected_blocks).max()),"initial":b["E2"].value,"final":b["E3"].value,"first_label":book_labels[0],"last_label":book_labels[-1],"template_label_mismatch_count":sum(a!=e for a,e in zip(book_labels,expected_labels))},
"dp_convergence":{"steps":steps.tolist(),"costs":costs.tolist(),"monotone_as_grid_refines":bool(np.all(np.diff(costs)<=1e-9))}}
assert milp["rows"] == dp["rows"] == 144
assert milp["max_balance_error"] < 1e-6 and milp["max_state_error"] < 1e-6
assert dp["max_balance_error"] < 1e-6 and dp["max_state_error"] < 1e-6
assert milp["simultaneous"] == dp["simultaneous"] == 0
assert report["result1"]["purchase_max_abs_diff"] < 1e-8
assert report["result1"]["block_max_abs_diff"] < 1e-8
assert report["result1"]["template_label_mismatch_count"] == 0
assert report["dp_convergence"]["monotone_as_grid_refines"]
assert abs(milp_summary["optimal_cost"] - milp["cost"]) < 1e-8
assert abs(milp_summary["optimal_purchase"] - milp["purchase"]) < 1e-8
(ROOT / "outputs" / "q1_audit_report.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(report,ensure_ascii=False,indent=2))
