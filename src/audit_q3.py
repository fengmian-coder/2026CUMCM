"""Independent saved-output and workbook checks for Q3."""
from pathlib import Path
import json,csv
import numpy as np
from openpyxl import load_workbook
from q2_data import load_q2_data
ROOT=Path(__file__).resolve().parents[1]

def main():
    out=ROOT/'outputs/q3';a=np.load(out/'main/schedules.npz');n=np.load(out/'main/natural.npz');data=load_q2_data()
    for k in ('purchase','charge','discharge','baseline','emergency','grid_fee'):
        np.testing.assert_array_equal(n[k][:,0],a[k][30:364,-1]);np.testing.assert_array_equal(n[k][:,1:],a[k][31:,:143])
    np.testing.assert_array_equal(a['purchase'][:,:36],a['baseline'][:,:36])
    np.testing.assert_array_equal(a['charge'][:,:36],a['baseline_charge'][:,:36])
    np.testing.assert_array_equal(a['discharge'][:,:36],a['baseline_discharge'][:,:36])
    state=a['energy'][:,1:]-a['energy'][:,:-1]-.9*a['charge']+a['discharge']/.9
    balance=a['purchase']+data.source_pv_kw/6+a['discharge']+a['emergency']-data.source_load_kw/6-a['charge']-a['surplus']
    assert np.abs(state).max()<1e-7 and np.abs(balance).max()<1e-7
    p=np.roll(data.price_yuan_per_kwh,-1)
    fee=p*(a['purchase']+.5*np.abs(a['purchase']-a['baseline']))
    np.testing.assert_allclose(fee,a['grid_fee'],atol=1e-9,rtol=0)
    events=list(csv.DictReader(open(out/'main/adjustment_events.csv',encoding='utf-8-sig')))
    assert set(int(e['hour']) for e in events)=={6,12}
    assert json.loads((out/'main/summary.json').read_text(encoding='utf-8'))['config']['hours']==[6,12]
    for e in events:
        d=(__import__('datetime').date.fromisoformat(e['date'])-data.dates[0]).days;t=int(e['hour'])*6
        assert abs(float(e['initial_energy_kwh'])-a['energy'][d,t])<1e-7
        assert float(e['expected_gain'])>=-1e-5
    w=load_workbook(ROOT/'materials/result3.xlsx',data_only=False)
    old=load_workbook(out/'result3_template_original.xlsx',data_only=False)
    for name in w.sheetnames:assert [c.value for c in w[name][1]]==[c.value for c in old[name][1]]
    max_cell=max_fee=max_bat=max_soc=0.
    for name,k in [('计划购电量','baseline'),('调整购电量','purchase')]:
        sh=w[name]
        for i,d in enumerate(range(31,365)):
            assert sh.cell(i+2,1).value==old[name].cell(i+2,1).value
            got=np.array([sh.cell(i+2,j).value for j in range(2,146)])
            max_cell=max(max_cell,float(np.abs(got-a[k][d]).max()))
            assert abs(sh.cell(i+2,146).value-a[k][d].sum())<1e-8
            expected=a['base_fee' if k=='baseline' else 'grid_fee'][d].sum()
            max_fee=max(max_fee,abs(sh.cell(i+2,147).value-expected))
    sh=w['充放电量']
    for i in range(334):
        for j in range(6):
            for col,k in [(3,'charge'),(4,'discharge')]:max_bat=max(max_bat,abs(sh.cell(2+i*6+j,col).value-n[k][i,j*24:(j+1)*24].sum()))
        max_soc=max(max_soc,abs(sh.cell(2+i*6,6).value-n['energy'][i,0]),abs(sh.cell(3+i*6,6).value-n['energy'][i,-1]))
    expected=[]
    def clock(t):return f'{t//6}:{t%6*10:02d}'
    for i,d in enumerate(data.dates[31:]):
        t=0
        while t<144:
            if n['emergency'][i,t]<=1e-7:t+=1;continue
            start=t
            while t<144 and n['emergency'][i,t]>1e-7:t+=1
            expected.append((str(d),f'{clock(start)}-{clock(t)}',float(n['emergency'][i,start:t].sum())))
    actual=[];day=None
    for row in w['紧急购电量'].iter_rows(min_row=2,max_col=3,values_only=True):
        if row[2] is None:continue
        if row[0] is not None:day=row[0].date().isoformat()
        actual.append((day,row[1],float(row[2])))
    assert len(actual)==len(expected)
    err=0.
    for got,exp in zip(actual,expected):assert got[:2]==exp[:2];err=max(err,abs(got[2]-exp[2]))
    assert max(max_cell,max_fee,max_bat,max_soc,err)<1e-8
    report=dict(plan_cell_error=max_cell,plan_fee_error=max_fee,battery_aggregation_error=max_bat,soc_error=max_soc,
                emergency_event_error=err,emergency_records=len(expected),state_residual=float(np.abs(state).max()),
                balance_residual=float(np.abs(balance).max()),initial_prefix_locked=True,stage_start_states_match=True,
                template_headers_dates_unchanged=True)
    (out/'audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(report)
if __name__=='__main__':main()
