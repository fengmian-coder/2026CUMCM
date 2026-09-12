"""Audit saved Q4 actions, actual-price settlement and both result templates."""
from pathlib import Path
import json
import numpy as np
from openpyxl import load_workbook
from q2_data import load_q2_data
R=Path(__file__).resolve().parents[1];O=R/'outputs/q4'
data=load_q2_data();price=np.load(O/'price_forecasts.npz')['actual'];reports=[]
for mode in [2,3]:
    with np.load(O/f'q{mode}/schedules.npz') as z:a={k:z[k] for k in z.files}
    with np.load(O/f'q{mode}/natural.npz') as z:n={k:z[k] for k in z.files}
    for k in ['purchase','charge','discharge','emergency','grid_fee']:
        np.testing.assert_array_equal(n[k][:,0],a[k][30:364,-1]);np.testing.assert_array_equal(n[k][:,1:],a[k][31:,:143])
    np.testing.assert_array_equal(a['purchase'][:,:36],a['baseline'][:,:36])
    expected=price*(a['purchase']+(.5*abs(a['purchase']-a['baseline']) if mode==3 else 0))
    assert abs(expected-a['grid_fee']).max()<1e-9
    assert abs(5*price*a['emergency']-a['emergency_fee']).max()<1e-9
    wb=load_workbook(R/f'materials/result4-{mode}.xlsx',data_only=True);old=load_workbook(O/f'result4-{mode}_template_original.xlsx')
    for name in wb.sheetnames:assert [c.value for c in wb[name][1]]==[c.value for c in old[name][1]]
    err=0.
    for name,k,fee in [('计划购电量','baseline','base_fee')]+([('调整购电量','purchase','grid_fee')] if mode==3 else []):
        sh=wb[name]
        for i,d in enumerate(range(31,365)):
            assert sh.cell(i+2,1).value==old[name].cell(i+2,1).value
            got=np.array([sh.cell(i+2,j).value for j in range(2,146)])
            err=max(err,float(abs(got-a[k][d]).max()),abs(sh.cell(i+2,146).value-a[k][d].sum()),abs(sh.cell(i+2,147).value-a[fee][d].sum()))
    sh=wb['充放电量']
    for i in range(334):
        for j in range(6):
            for col,k in [(3,'charge'),(4,'discharge')]:err=max(err,abs(sh.cell(2+i*6+j,col).value-n[k][i,j*24:(j+1)*24].sum()))
        err=max(err,abs(sh.cell(2+i*6,6).value-n['energy'][i,0]),abs(sh.cell(3+i*6,6).value-n['energy'][i,-1]))
    expected=[]
    clock=lambda t:f'{t//6}:{t%6*10:02d}'
    for i,dt in enumerate(data.dates[31:]):
        t=0
        while t<144:
            if n['emergency'][i,t]<=1e-7:t+=1;continue
            start=t
            while t<144 and n['emergency'][i,t]>1e-7:t+=1
            expected.append((str(dt),f'{clock(start)}-{clock(t)}',float(n['emergency'][i,start:t].sum())))
    actual=[];day=None
    for row in wb['紧急购电量'].iter_rows(min_row=2,max_col=3,values_only=True):
        if row[2] is None:continue
        if row[0] is not None:day=row[0].date().isoformat()
        actual.append((day,row[1],float(row[2])))
    assert len(actual)==len(expected)
    for g,e in zip(actual,expected):assert g[:2]==e[:2];err=max(err,abs(g[2]-e[2]))
    assert err<1e-7
    wb.close();old.close();reports.append(dict(mode=mode,max_workbook_error=err,emergency_records=len(expected),template_headers_dates_unchanged=True,natural_mapping_exact=True,actual_price_settlement_verified=True))
(O/'workbook_audit.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8');print(reports)
