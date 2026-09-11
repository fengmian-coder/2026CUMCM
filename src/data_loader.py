"""Strict readers for the CUMCM microgrid source workbooks.

The Excel files are read-only. Question 1 uses left-endpoint constant power
and a periodic boundary assumption: the final midnight sample is moved to
the start of the natural-day solve. Other daily matrices retain source order.
"""

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import re

import numpy as np
import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MATERIALS = PROJECT_ROOT / "materials"
STEPS_PER_DAY = 144
STEP_MINUTES = 10
DT_HOURS = STEP_MINUTES / 60
EXPECTED_ENDPOINT_MINUTES = np.arange(10, 24 * 60 + 1, STEP_MINUTES)


@dataclass(frozen=True)
class DayData:
    interval_start_minutes: np.ndarray
    interval_end_minutes: np.ndarray
    interval_labels: tuple[str, ...]
    price_yuan_per_kwh: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray

    @property
    def load_kwh(self):
        return self.load_kw * DT_HOURS

    @property
    def pv_kwh(self):
        return self.pv_kw * DT_HOURS


def _endpoint_minutes(value) -> int:
    if isinstance(value, time):
        minutes = value.hour * 60 + value.minute
        return 24 * 60 if minutes == 0 else minutes
    if isinstance(value, datetime):
        minutes = value.hour * 60 + value.minute
        return 24 * 60 if minutes == 0 else minutes
    if isinstance(value, (int, float)):
        minutes = int(round(float(value) * 24 * 60))
        return 24 * 60 if minutes == 0 else minutes
    text = str(value).strip().replace("：", ":")
    next_day = "+1" in text
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?:\+1)?", text)
    if not match:
        raise ValueError(f"Unrecognized Excel time value: {value!r}")
    hour, minute = map(int, match.groups())
    result = hour * 60 + minute
    if next_day and result == 0:
        result = 24 * 60
    return result


def _clock(minutes: int) -> str:
    if minutes == 24 * 60:
        return "24:00"
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def _interval_labels(endpoints: np.ndarray) -> tuple[str, ...]:
    return tuple(
        f"{_clock(int(end - STEP_MINUTES))}-{_clock(int(end))}"
        for end in endpoints
    )


def _assert_finite_nonnegative(
    name: str, values: np.ndarray, strictly_positive=False, expected_width=STEPS_PER_DAY
):
    if values.shape[-1] != expected_width:
        raise ValueError(f"{name}: expected {expected_width} values, got {values.shape[-1]}")
    if not np.isfinite(values).all():
        raise ValueError(f"{name}: contains missing or non-finite values")
    if strictly_positive and np.any(values <= 0):
        raise ValueError(f"{name}: contains a non-positive value")
    if not strictly_positive and np.any(values < 0):
        raise ValueError(f"{name}: contains a negative value")


def read_attachment1(path: Path | None = None) -> DayData:
    path = path or MATERIALS / "附件1.xlsx"
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.values)
    if rows[0][:4] != ("时间", "电价", "小区负载", "光伏发电预测功率"):
        raise ValueError(f"附件1表头发生变化: {rows[0][:4]}")
    if len(rows) != STEPS_PER_DAY + 1:
        raise ValueError(f"附件1应有145行（含表头），实际为{len(rows)}行")

    endpoints = np.asarray([_endpoint_minutes(row[0]) for row in rows[1:]], dtype=int)
    if not np.array_equal(endpoints, EXPECTED_ENDPOINT_MINUTES):
        raise ValueError("附件1时间轴不是00:10至24:00的连续10分钟右端点")
    data = np.asarray([[row[1], row[2], row[3]] for row in rows[1:]], dtype=float)
    # Explicit Q1 periodic boundary assumption; preserve every raw value.
    data = np.concatenate((data[-1:], data[:-1]), axis=0)
    price, load, pv = data.T
    _assert_finite_nonnegative("附件1电价", price, strictly_positive=True)
    _assert_finite_nonnegative("附件1负载", load)
    _assert_finite_nonnegative("附件1光伏", pv)
    return DayData(
        interval_start_minutes=endpoints - STEP_MINUTES,
        interval_end_minutes=endpoints,
        interval_labels=_interval_labels(endpoints),
        price_yuan_per_kwh=price,
        load_kw=load,
        pv_kw=pv,
    )


def _read_daily_matrix(path: Path, sheet_name: str):
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet_name]
    rows = list(sheet.values)
    endpoints = np.asarray([_endpoint_minutes(v) for v in rows[0][1:]], dtype=int)
    if not np.array_equal(endpoints, EXPECTED_ENDPOINT_MINUTES):
        raise ValueError(f"{path.name}/{sheet_name}: 时间轴不连续或端点解释不一致")
    dates = tuple(row[0].date() if isinstance(row[0], datetime) else row[0] for row in rows[1:])
    expected_dates = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(365))
    if dates != expected_dates:
        raise ValueError(f"{path.name}/{sheet_name}: 日期不是2025全年连续日期")
    values = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    _assert_finite_nonnegative(f"{path.name}/{sheet_name}", values)
    return dates, values


def read_attachment2(path: Path | None = None):
    path = path or MATERIALS / "附件2.xlsx"
    dates_l, load = _read_daily_matrix(path, "小区负载")
    dates_p, pv = _read_daily_matrix(path, "光伏发电实际功率")
    if dates_l != dates_p:
        raise ValueError("附件2负载与光伏日期不一致")
    return dates_l, load, pv


def read_attachment4(path: Path | None = None):
    path = path or MATERIALS / "附件4.xlsx"
    dates, price = _read_daily_matrix(path, "Sheet1")
    _assert_finite_nonnegative("附件4电价", price, strictly_positive=True)
    return dates, price


def read_attachment3(path: Path | None = None):
    path = path or MATERIALS / "附件3.xlsx"
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.values)
    expected_header = ("日期", "预报时刻") + tuple(f"预报{i}小时" for i in range(1, 25))
    if tuple(rows[0]) != expected_header:
        raise ValueError("附件3表头发生变化")
    if len(rows) != 365 * 4 + 1:
        raise ValueError("附件3记录数不是365天×4个发布时间")

    forecasts = []
    current_date = None
    expected_issue_hours = (0, 6, 12, 18)
    for index, row in enumerate(rows[1:]):
        if row[0] not in (None, ""):
            current_date = datetime.strptime(str(row[0]), "%Y-%m-%d").date()
        if current_date is None:
            raise ValueError("附件3首条记录缺少日期")
        issue_hour = int(str(row[1]).split(":", 1)[0])
        if issue_hour != expected_issue_hours[index % 4]:
            raise ValueError(f"附件3第{index + 2}行预报时刻顺序异常")
        values = np.asarray(row[2:], dtype=float)
        _assert_finite_nonnegative("附件3光伏预报", values, expected_width=24)
        forecasts.append((datetime.combine(current_date, time(issue_hour)), values))
    return tuple(forecasts)


def validate_all_sources():
    a1 = read_attachment1()
    dates, load, pv = read_attachment2()
    forecasts = read_attachment3()
    price_dates, price = read_attachment4()
    if dates != price_dates:
        raise ValueError("附件2与附件4日期不一致")
    return {
        "attachment1_intervals": len(a1.interval_labels),
        "attachment1_first_interval": a1.interval_labels[0],
        "attachment1_last_interval": a1.interval_labels[-1],
        "attachment2_days": len(dates),
        "attachment2_shape": load.shape,
        "attachment3_forecasts": len(forecasts),
        "attachment4_shape": price.shape,
    }


if __name__ == "__main__":
    for key, value in validate_all_sources().items():
        print(f"{key}: {value}")
