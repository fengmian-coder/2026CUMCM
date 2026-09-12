"""Create the implemented Q3 method report with native Word equations."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'outputs/q3'
def read(path):return json.loads((OUT/path).read_text(encoding='utf-8'))
main=read('main/summary.json');cases={v['case']:v for v in read('comparison.json')};metrics=read('forecast_metrics_w28.json');audit=read('audit.json')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
    s=doc.styles[name];s.font.name='Times New Roman';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体');s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(11 if name=='Normal' else 23 if name=='Title' else 16 if name=='Heading 1' else 12)
    s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.18
doc.styles['Normal'].paragraph_format.first_line_indent=Pt(22)
for name in ['Heading 1','Heading 2']:doc.styles[name].paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.text='第三问建模与求解过程';header.style='Normal';header.paragraph_format.first_line_indent=0
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(t,bold=False):
    x=doc.add_paragraph();r=x.add_run(t);r.bold=bold;return x
def h(t):doc.add_heading(t,level=1)
def subhead(t):doc.add_heading(t,level=2)
def page():doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for j,v in enumerate(headers):t.rows[0].cells[j].text=str(v)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):cells[j].text=str(v)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
    pr.append(b)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8ED' if i==0 else 'FFFFFF');tcp.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcp.append(margins)
            for par in c.paragraphs:
                par.paragraph_format.first_line_indent=0;par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.05
                for r in par.runs:r.font.size=Pt(10);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=0
    return t

def txt(s):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
def sub(base,index):
    x=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(txt(base));i=OxmlElement('m:sub');i.append(txt(index));x.extend([e,i]);return x
def frac(a,b):
    x=OxmlElement('m:f');n=OxmlElement('m:num');n.append(txt(a));d=OxmlElement('m:den');d.append(txt(b));x.extend([n,d]);return x
def eq(*parts):
    par=doc.add_paragraph();par.paragraph_format.first_line_indent=0;par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    m=OxmlElement('m:oMath');
    for item in parts:m.append(txt(item) if isinstance(item,str) else item)
    par._p.append(m)
def picture(file,width=6.65):
    par=doc.add_paragraph();par.paragraph_format.first_line_indent=0;par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.add_run().add_picture(str(ROOT/'figures'/file),width=Inches(width))

doc.add_paragraph('第三问建模与求解过程',style='Title')
p('基于历史预测择优与调整价值判别的多阶段滚动优化',True)
p('本报告供论文手和建模手复核，记录本次实际运行的模型、参数、求解流程与数值结果。实现依据为《第三问-2》及随后确认的修订意见。文中的假设与原题要求分别说明，结果取自本地附件与程序输出。')
p(f"四时点主方案在2025年2月1日至12月31日自然日窗口内，总费用为{main['natural_totals']['total_cost_yuan']}元。相对仅使用00:00预报的第三问对照，费用下降125025.82453696057元。加入18:00更新并未进一步降低本次全年费用，不能将预报次数越多视为必然更优。")
h('一 时间轴与决策边界')
p('沿用第二问已确认的建模假设：每天00:00制定的144段计划从当天00:10开始，至次日00:10结束。每个整点发布的新信息只影响下一个10分钟边界以后的动作；这属于与模板协调的实施假设，并非原题明示的通信延迟。')
table(['发布时间','允许重新优化的序号','对应区间'],[['00:00','1—144','00:10—次日00:10'],['06:00','37—144','06:10—次日00:10'],['12:00','73—144','12:10—次日00:10'],['18:00','109—144','18:10—次日00:10']],[1.1,1.9,3.7])
p('每次重算全部剩余时段，而非只重算接下来的6小时。此前已执行及当前整点到下一10分钟的动作被冻结。下一轮以当前有效方案继续递推，因此既不会回滚已执行动作，也无需在6小时边界强制回到原SOC。')
p('仿真从1月1日开始连续运行。1月1日00:00—00:10储能静置，00:10初始电量为6000 kWh。之后每天的初始状态承接前一天最终方案末端，不从2月重新初始化。')

page();h('二 数据读取与三类预测')
p('附件1提供分时电价，附件2提供实际负荷和光伏功率，附件3提供1460批整点光伏预报。预报1小时对应发布时间后一小时；每批共有24个未来整点。原始工作簿只读，数值不先舍入。')
p('Q候选直接取当天00:00生成的第二问光伏预测在剩余区间的切片。负载始终使用当天第二问负载预测；本次没有另建日内负载更新模型。')
p('A候选由当前已知光伏实测值与附件3本批未来整点预报直接线性插值得到。当前观测的可及时获取是一项信息假设；1月1日零点无实测记录时采用附件1冷启动先验，不拿未来00:10实测值冒充。')
p('F候选保留Q的10分钟细分形态，通过整点锚点偏差修正。设发布时间为r，j表示距发布时间的整点数，Q和A为同一目标时刻的两种预测。')
eq(sub('Z','j'),' = w ',sub('Q','j'),' + (1 − w) ',sub('A','j'),'， j ≥ 1')
eq(sub('Z','0'),' = ',sub('P','r,obs'),'， ',sub('Δ','j'),' = ',sub('Z','j'),' − ',sub('Q','j'))
p('在相邻整点之间，对修正量Δ线性插值，叠加到Q原有曲线上。整点本身直接取融合锚点。首个小时起点采用当前观测，因此F在权重等于1时仍可能与未校正的Q在首小时不同。')
eq(sub('F','k'),' = max{0， ',sub('Q','k'),' + interp(',sub('Δ','j'),') }')
p('未发现原题给出可直接使用的光伏装机容量，因此本次不使用虚构容量截断，也不把储能5000 kW功率上限套在光伏上。不依据全年实测数据事后强制夜间置零，只保留非负约束，避免引入未来信息。')
p('程序优化变量统一使用每10分钟的电量，单位kWh。预测功率乘以1/6小时转为电量。显示精度与计算精度分开，结果文件保留浮点精度。')
eq('Δt = ',frac('1','6'),' h， ',sub('Q','k,PV'),' = ',sub('P','k,PV'),' Δt')

page();h('三 历史择优与不确定性场景')
p('每个发布时间分别使用最近28个已经完整观测的历史同发布时间批次。前期不足28批时使用已有批次；无历史批次的首日固定选Q。权重在0至1之间按0.01步长搜索，在最终10分钟曲线上计算误差，不仅比较整点。')
p('先选历史MAE最低的融合权重，精确并列时依次取较低RMSE及较小权重。再比较Q、A及最佳F：MAE落在最小值的2%范围内时取RMSE较小者，完全并列时优先F。本次采用这一确定规则，没有加入难以界定的“无显著差异”判断，也没有临时使用全年调度费用打破并列。')
p('历史数据只参与当时的选权重和选模型。评估报告另按2—12月真实的逐日预测结果计算误差，因此历史训练误差不冒充最终预测精度。14日和42日窗口在完整独立仿真中用于敏感性检验。')
table(['发布时间','Q的MAE','A的MAE','择优后MAE'],[[h,f"{metrics[h]['Q']['mae_kw']:.4f}",f"{metrics[h]['A']['mae_kw']:.4f}",f"{metrics[h]['selected']['mae_kw']:.4f}"] for h in ['0','6','12','18']],[1.25,1.8,1.8,1.8])
p('表中MAE单位为kW。各发布时间评估的剩余区间长度不同，不能直接把18:00较低的MAE解释为其预报技术最优；夜间光伏功率本身接近零。')
p('每次优化生成30个联合负荷和光伏残差场景。历史窗口60日，年龄权重半衰期30日，同星期权重乘3；抽取同一历史日期的整条负荷与光伏残差，保留日内及源荷之间的联合变化。场景等概率，随机种子固定以便复现。')
p('光伏残差与当前候选来源、发布时间及融合权重相匹配：在历史输入上重建该候选，再与历史实际值比较。没有将第二问Q的残差直接套用到A或F。零点及各日内时点仅使用此前日期已完整观测的残差批次，首日使用单一预测场景。')
p('改变2025年6月21日12:10以后实测数据和更晚发布的预报，不改变当日12:00及此前的预测、权重和来源选择，信息边界扰动核验通过。')

page();h('四 费用结算与优化目标')
p('本次主模型明确采用“减购退回对应原购电费，另付50%违约费”的解释，增加部分按1.5倍电价结算。该退费解释列为建模假设，并独立运行不退费版本检验影响。各次调整始终以当天00:00基准G⁰为费用参照，而非逐次累计收费。')
p('以下G、C、D、B均为单时段电量，p为元/kWh。调整后的完整外网结算费用为：')
eq(sub('c','k,grid'),' = ',sub('p','k'),sub('G','k'),' + 0.5 ',sub('p','k'),' |',sub('G','k'),' − ',sub('G','k,0'),'|')
p('这个公式已经包含调整后的购电结算，不能再次叠加完整原计划费。报表中的“调整净费用”等于完整结算费减去基准计划费，个别时段可能为负。')
p('不退费敏感性版本使用：')
eq(sub('c','k,grid'),' = ',sub('p','k'),sub('G','k,0'),' + 0.5 ',sub('p','k'),'(',sub('G','k,0'),' − ',sub('G','k'),')⁺ + 1.5 ',sub('p','k'),'(',sub('G','k'),' − ',sub('G','k,0'),')⁺')
p('对场景s，紧急购电费用Lₛ为剩余区间内5倍电价紧急购电量之和。CVaR仅作用于紧急购电费用，参数与第二问一致：α=0.90、λ=0.05。确定性购电费不再额外乘风险权重。')
eq(sub('L','s'),' = ∑ ',sub('5p','k'),sub('B','s,k'))
eq('min J = ∑ ',sub('c','k,grid'),' + ∑ ',sub('π','s'),sub('L','s'),' + λ CVaR(',sub('L','s'),') − v ',sub('E','end'))
eq('CVaR = ζ + ',frac('1','1 − α'),' ∑ ',sub('π','s'),sub('u','s'),'， ',sub('u','s'),' ≥ ',sub('L','s'),' − ζ， ',sub('u','s'),' ≥ 0')
p('终端价值v=0.6895775元/kWh，终端时刻始终为次日00:10。该价值项用于优化决策，不计入事后现金购电费用。终端电量仅受1200—10800 kWh范围约束，不强制等于第二问曾经求得的某个数值。')

page();h('五 约束与滚动求解')
p('所有场景共享计划购电和充放电动作，场景紧急购电与剩余电量分别处理供需不足和过剩。剩余量是等效供电过剩的核算变量，可能来自计划购电或光伏，不能全部称为弃光。')
eq(sub('Q','s,k,PV'),' + ',sub('G','k'),' + ',sub('D','k'),' + ',sub('B','s,k'),' = ',sub('Q','s,k,L'),' + ',sub('C','k'),' + ',sub('W','s,k'))
eq(sub('E','k+1'),' = ',sub('E','k'),' + 0.9 ',sub('C','k'),' − ',frac('Dₖ','0.9'))
eq('1200 ≤ ',sub('E','k'),' ≤ 10800， 0 ≤ ',sub('C','k'),'，',sub('D','k'),' ≤ ',frac('5000','6'),' kWh')
p('每轮从当前有效方案在r+10分钟处的储电量开始。先以当前动作和最新相同场景计算keep，再放开剩余动作计算adjust。keep与adjust使用相同的费用规则、风险参数与终端价值。')
eq('ΔJ = ',sub('J','keep'),' − ',sub('J','adjust'))
eq('调整条件：ΔJ > max{ ε max(',sub('C','keep,cash'),'，1元)，10⁻⁷元 }')
p('主阈值ε=0.001。比例基数采用正的预期现金费用，排除终端价值项，避免优化目标为负或接近零时收益率失真。若未超过阈值，则保留当前方案；否则仅替换可调整集合内的购电、充放电及未来SOC。')
p('线性规划允许不改变当前动作，故adjust最优目标不应高于keep。本次每次求解均检查该关系，同时独立重算求解器目标。采用SciPy HiGHS求解；若LP解出现同段同时充放电，则自动加入充放电状态二元变量重新求MILP。实际主方案没有触发该回退，所有已保存时段均未同时充放电。')
p('在18:00更新后，次日00:00—00:10也保持锁定；次日零点报告前143段执行后的电量，次日新优化使用再执行末段后的00:10电量。实际源荷只用于事后紧急购电和剩余量结算，不用于提前挑选当天动作。')

page();h('六 主方案结果与预报时点价值')
p('以下均按相同自然日窗口统计：2025年2月1日00:00至2026年1月1日00:00。终端价值不计入现金总费用，各对照均从1月1日独立递推，不复用主方案的SOC轨迹。')
table(['主方案指标','计算结果'],[['基准购电费 元',main['natural_totals']['base_plan_fee_yuan']],['调整净费用 元',main['natural_totals']['adjustment_fee_yuan']],['紧急购电费 元',main['natural_totals']['emergency_fee_yuan']],['总费用 元',main['natural_totals']['total_cost_yuan']],['紧急购电量 kWh',main['natural_totals']['emergency_kwh']]], [2.7,3.9])
table(['预报更新组合','自然日总费用 元'],[[label,cases[k]['natural_totals']['total_cost_yuan']] for k,label in [('only_0','仅00:00'),('updates_06','00:00及06:00'),('updates_0612','00:00及06:00及12:00'),('main','四时点全部使用')]], [2.7,3.9])
p('四时点方案相对只用00:00节省125025.82453696057元，约0.8610%。在已加入06:00和12:00的基础上再加入18:00，总费用增加2259.214770646766元。本次18:00更新没有体现进一步省钱的价值，但这一结果是样本期回测结论，不能当作未来所有日期的保证。')
p('18:00主方案触发325次调整，06:00和12:00均触发334次。触发次数反映当时预计收益超过阈值，不等于实际收益次数。预测分布与真实值有偏差，跨日储电量也会影响后续成本，因此预期改善不保证事后费用下降。')
p('当前result3保存事先定义的四时点主方案；00:00加06:00加12:00方案作为本次更经济的对照保留，未依据全年已知结果偷偷替换主方案。四种组合并未穷尽全部8种更新子集。')

page();h('七 对照结果图与敏感性')
picture('q3_update_cost_comparison.png',6.5)
table(['试验设置','总费用 元'],[[label,cases[k]['natural_totals']['total_cost_yuan']] for k,label in [('window14','历史窗口14日'),('window42','历史窗口42日'),('threshold0','调整阈值0'),('threshold005','调整阈值0.5%'),('no_refund','不退费口径')]], [2.7,3.9])
p('上述窗口和阈值试验均使用四个发布时间，且每组独立运行全年。主参数28日、0.1%在运行前确定；报告如实呈现其他参数结果，不把完整回测期同时作为参数选择后的独立测试集。')
p('不退费口径总费用为14448478.051924346元，说明退费解释确实影响结论；它的紧急购电量更低，但不能据此说整体经济性更好。应同时比较完整现金费用和供电风险。')

page();h('八 预测结果与结果表填写')
picture('q3_forecast_mae.png',6.45)
table(['发布时间','选择Q的天数','选择A的天数','选择F的天数'],[[h,*[metrics[h]['selection_counts'][k] for k in ('Q','A','F')]] for h in ('0','6','12','18')],[1.4,1.75,1.75,1.75])
p('result3的计划购电量页保存00:00基准G⁰；调整购电量页保存最终有效购电量G，不是有正负号的差值。两页均按各自144段相同时窗填写。调整购电量页的全天购电费是完整调整后外网结算费，已含调整影响，但不含紧急购电。不能将两页全天购电费直接相加。')
p('充放电页按自然日六个4小时区间保存最终动作，零点及24:00电量取正确边界。紧急购电页按自然日合并连续正缺口区间，无紧急购电日期不另列零记录。共334个计划日、2004行充放电汇总和3206段紧急购电记录。')
p('计划表行合计覆盖2月1日00:10至次年1月1日00:10；自然日统计覆盖2月1日00:00至次年1月1日00:00。两者窗口不同，不混加。主方案计划窗口总费用为14395421.07897714元。')

page();h('九 核验结论与复现文件')
table(['检查项目','结果'],[['购电逐格写入最大误差',audit['plan_cell_error']],['SOC写入最大误差',audit['soc_error']],['费用行合计最大误差 元',audit['plan_fee_error']],['状态递推最大残差 kWh',audit['state_residual']],['电量平衡最大残差 kWh',audit['balance_residual']],['跨日SOC衔接误差 kWh',main['audit']['continuity_residual_kwh']],['同时充放电时段数',main['audit']['simultaneous_slots']]], [3.6,3.0])
p('除浮点运算顺序造成的约10⁻¹¹至10⁻¹²级误差外，购电格映射、零点状态、执行冻结边界及费用结算关系均通过检查。模板表头与日期逐项保持原样；检查也覆盖每一段紧急购电的日期、起止时刻和电量。')
p('正式运行入口为src/refresh_q3.py，依次生成14日、28日和42日预测，再运行主方案、预报时点对照、退费口径及阈值敏感性。预测、优化和仿真分别由q3_forecast.py、q3_solver.py、q3_run.py实现。求解依赖NumPy和SciPy，固定随机种子保证可复现。')
p('outputs/q3/main保存主方案逐时执行表execution.csv、计划与最终动作数组schedules.npz、自然日报告数组natural.npz以及每次调整的keep/adjust比较adjustment_events.csv。outputs/q3/comparison.json保存九组试验汇总。forecast_selection_w28.csv记录每天各时点的权重、来源与历史评分。')
p('audit_q3.py检查表格和物理约束；audit_q3_forecast.py检查未来数据扰动不改变此前预测。src/fill_result3.mjs使用工作区电子表格工具填写materials/result3.xlsx。原始附件1至3未修改，所有输出仅来自本地数据及明确记载的冷启动假设。')
subhead('写入论文时必须保留的限定')
p('退费解释、10分钟生效边界、当前整点观测可及时获取均属于建模假设。光伏择优不代表所有时点融合必然最好。所有比较是本数据集上的滚动回测；终端SOC略有差别，费用对比保持同一终端价值政策，但不是强制相同终端电量下的全局最优证明。')
p('第三问相对第二问总费用下降320038.62830308266元，同时改变了零点光伏输入和日内更新机制，不能将全部差额归因于日内预报。单独评价日内更新，应采用第三问仅00:00方案作为对照。')
for root in (doc._element, doc.styles.element):
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
path=OUT/'第三问建模与求解过程.docx';doc.save(path);print(path)
