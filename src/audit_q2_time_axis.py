"""Independent timestamp, information-boundary, and template-preservation checks."""
from pathlib import Path
from dataclasses import replace
import json
import numpy as np
from openpyxl import load_workbook
from q2_data import load_q2_data
from q2_forecast import rolling_forecasts

ROOT=Path(__file__).resolve().parents[1]
def main():
    out=ROOT/'outputs/q2';data=load_q2_data()
    p=np.load(out/'q2_schedules.npz');n=np.load(out/'natural_reporting.npz')
    for key in ('purchase','charge','discharge','emergency','surplus'):
        np.testing.assert_array_equal(n[key][:,0],p[key][30:364,-1])
        np.testing.assert_array_equal(n[key][:,1:],p[key][31:,:143])
    np.testing.assert_array_equal(n['energy'][:,0],p['energy'][30:364,143])
    np.testing.assert_array_equal(n['energy'][:,1:],p['energy'][31:,:144])
    state=n['energy'][:,1:]-n['energy'][:,:-1]-.9*n['charge']+n['discharge']/.9
    assert np.max(np.abs(state))<1e-7
    # Alter only as-yet-unobserved samples. Earlier/current forecasts must not change.
    cutoff=171
    L=data.source_load_kw.copy();P=data.source_pv_kw.copy()
    L[cutoff:]+=4321;P[cutoff:]+=1234
    changed=replace(data,source_load_kw=L,source_pv_kw=P)
    lh,ph=rolling_forecasts(changed)
    original=np.load(out/'rolling_forecasts.npz')
    np.testing.assert_array_equal(lh[:cutoff+1],original['load_kw'][:cutoff+1])
    np.testing.assert_array_equal(ph[:cutoff+1],original['pv_kw'][:cutoff+1])
    # Template label/date preservation compared with the archived pre-edit file.
    old=load_workbook(ROOT/'outputs/q2_legacy_natural/result2.xlsx',data_only=False)
    new=load_workbook(ROOT/'materials/result2.xlsx',data_only=False)
    for name in old.sheetnames:
        assert [c.value for c in old[name][1]]==[c.value for c in new[name][1]]
    for r in range(2,336):
        assert old['计划购电量'].cell(r,1).value==new['计划购电量'].cell(r,1).value
    report={'natural_execution_mapping':'exact','zero_and_0010_soc_mapping':'exact',
            'natural_state_residual_kwh':float(np.abs(state).max()),
            'future_data_perturbation_date':str(data.dates[cutoff]),
            'earlier_and_current_forecasts_unchanged':True,'template_headers_and_plan_dates_unchanged':True,
            'first_plan_initial_energy_kwh':float(p['energy'][0,0]),
            'last_plan_first_decision_date':'2025-12-31','last_plan_interval':'2026-01-01 00:00-00:10',
            'last_plan_purchase_kwh':float(p['purchase'][-1,-1])}
    (out/'time_axis_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
