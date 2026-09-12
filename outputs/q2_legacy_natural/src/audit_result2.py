"""Independent consistency checks for the filled result2.xlsx workbook."""

from datetime import datetime
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def as_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def main():
    out = ROOT / "outputs" / "q2"
    saved = np.load(out / "q2_schedules.npz")
    extension = json.loads((out / "q2_extension_2026-01-01.json").read_text("utf-8"))
    # Normal mode gives O(1) random cell access; read-only worksheets rescan XML
    # for every cell and are prohibitively slow for this cross-check.
    wb = load_workbook(ROOT / "materials" / "result2.xlsx", data_only=False, read_only=False)
    plan = wb["计划购电量"]
    battery = wb["充放电量"]
    emergency = wb["紧急购电量"]

    purchase = saved["purchase"]
    charge = saved["charge"]
    discharge = saved["discharge"]
    energy = saved["energy"]
    emergency_values = saved["emergency"]
    max_plan_error = max_total_error = 0.0
    for offset, d in enumerate(range(31, 365)):
        row = offset + 2
        tail = purchase[d + 1, 0] if d < 364 else extension["first_interval"]["purchase_kwh"]
        expected = np.r_[purchase[d, 1:], tail]
        actual = np.asarray([plan.cell(row, col).value for col in range(2, 146)], dtype=float)
        max_plan_error = max(max_plan_error, float(np.max(np.abs(actual - expected))))
        max_total_error = max(max_total_error, abs(float(plan.cell(row, 146).value) - float(expected.sum())))

    max_charge_error = max_discharge_error = max_energy_error = 0.0
    for offset, d in enumerate(range(31, 365)):
        base = 2 + offset * 6
        for block in range(6):
            expected_c = float(charge[d, block*24:(block+1)*24].sum())
            expected_u = float(discharge[d, block*24:(block+1)*24].sum())
            max_charge_error = max(max_charge_error, abs(float(battery.cell(base+block, 3).value)-expected_c))
            max_discharge_error = max(max_discharge_error, abs(float(battery.cell(base+block, 4).value)-expected_u))
        max_energy_error = max(
            max_energy_error,
            abs(float(battery.cell(base, 6).value)-float(energy[d, 0])),
            abs(float(battery.cell(base+1, 6).value)-float(energy[d, -1])),
        )

    emergency_total = 0.0
    emergency_rows = 0
    for row in emergency.iter_rows(min_row=2, max_col=3, values_only=True):
        if row[2] is None:
            continue
        emergency_total += float(row[2])
        emergency_rows += 1
    expected_emergency = float(emergency_values[31:].sum())
    report = {
        "plan_dates": [as_date(plan.cell(2, 1).value), as_date(plan.cell(335, 1).value)],
        "plan_mapping_max_abs_error_kwh": max_plan_error,
        "plan_total_max_abs_error_kwh": max_total_error,
        "battery_charge_max_abs_error_kwh": max_charge_error,
        "battery_discharge_max_abs_error_kwh": max_discharge_error,
        "battery_soc_max_abs_error_kwh": max_energy_error,
        "emergency_rows": emergency_rows,
        "emergency_total_kwh": emergency_total,
        "expected_emergency_total_kwh": expected_emergency,
        "emergency_total_abs_error_kwh": abs(emergency_total-expected_emergency),
        "template_first_interval": plan.cell(1, 2).value,
        "template_last_interval": plan.cell(1, 145).value,
    }
    assert report["plan_dates"] == ["2025-02-01", "2025-12-31"]
    assert max_plan_error < 1e-9 and max_total_error < 1e-8
    assert max_charge_error < 1e-8 and max_discharge_error < 1e-8
    assert max_energy_error < 1e-8
    assert report["emergency_total_abs_error_kwh"] < 1e-8
    assert report["template_first_interval"] == "0:10-0:20"
    assert report["template_last_interval"] == "0:00-0:10+1"
    (out / "result2_audit_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
